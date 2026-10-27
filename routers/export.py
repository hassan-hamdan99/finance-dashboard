import io
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from analytics.insights import get_transactions_df, monthly_summary, pivot_by_category_month

router = APIRouter(prefix="/api/export", tags=["export"])


def _build_workbook(df, months: int) -> Workbook:
    """Build an Excel workbook with 3 sheets."""
    wb = Workbook()

    # ── Sheet 1: Gastos mensais ──────────────────────────────────
    ws1 = wb.active
    ws1.title = "Gastos mensais"

    headers = [
        "DATA", "CLASSIFICACAO", "PROVEDOR", "DESCRICAO", "CONTA",
        "PAGAMENTO", "VALOR", "CATEGORIA", "SUBCATEGORIA", "SUBDIVISAO", "RECORRENCIA",
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    if not df.empty:
        sorted_df = df.sort_values("date", ascending=False)
        for row_idx, (_, row) in enumerate(sorted_df.iterrows(), 2):
            date_val = str(row.get("date", ""))[:10]
            ws1.cell(row=row_idx, column=1, value=date_val).border = thin_border
            ws1.cell(row=row_idx, column=2, value="Gasto" if row.get("amount", 0) < 0 else "Entrada").border = thin_border
            ws1.cell(row=row_idx, column=3, value=row.get("vendor", "")).border = thin_border
            ws1.cell(row=row_idx, column=4, value=row.get("description", "")).border = thin_border
            ws1.cell(row=row_idx, column=5, value=row.get("bank", "")).border = thin_border
            ws1.cell(row=row_idx, column=6, value=row.get("payment_method", "")).border = thin_border
            val_cell = ws1.cell(row=row_idx, column=7, value=round(row.get("amount", 0), 2))
            val_cell.number_format = '#,##0.00'
            val_cell.border = thin_border
            ws1.cell(row=row_idx, column=8, value=row.get("category_name", "")).border = thin_border
            ws1.cell(row=row_idx, column=9, value=row.get("subcategory", "")).border = thin_border
            ws1.cell(row=row_idx, column=10, value=row.get("subdivision", "")).border = thin_border
            ws1.cell(row=row_idx, column=11, value="Sim" if row.get("is_recurring") else "Nao").border = thin_border

    # Auto-width columns
    for col in ws1.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws1.column_dimensions[col_letter].width = min(max_length + 2, 40)

    # ── Sheet 2: Entradas e Saidas ───────────────────────────────
    ws2 = wb.create_sheet("Entradas e Saidas")
    summary = monthly_summary(df)

    if not summary.empty:
        month_cols = list(summary["month"])
        # Header row
        ws2.cell(row=1, column=1, value="").fill = header_fill
        ws2.cell(row=1, column=1).font = header_font
        for ci, m in enumerate(month_cols, 2):
            cell = ws2.cell(row=1, column=ci, value=m)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        # Add total column
        total_cell = ws2.cell(row=1, column=len(month_cols) + 2, value="TOTAL")
        total_cell.fill = header_fill
        total_cell.font = header_font

        row_labels = ["Entradas", "Saidas", "Resultado final"]
        row_keys = ["receitas", "gastos", "saldo"]
        for ri, (label, key) in enumerate(zip(row_labels, row_keys), 2):
            ws2.cell(row=ri, column=1, value=label).font = Font(bold=True)
            row_total = 0.0
            for ci, m in enumerate(month_cols, 2):
                val = summary[summary["month"] == m][key].values
                v = round(float(val[0]), 2) if len(val) > 0 else 0.0
                cell = ws2.cell(row=ri, column=ci, value=v)
                cell.number_format = '#,##0.00'
                row_total += v
            total_val_cell = ws2.cell(row=ri, column=len(month_cols) + 2, value=round(row_total, 2))
            total_val_cell.number_format = '#,##0.00'
            total_val_cell.font = Font(bold=True)

    # ── Sheet 3: Tabela Dinamica ─────────────────────────────────
    ws3 = wb.create_sheet("Tabela Dinamica")
    pivot = pivot_by_category_month(df, months=months)

    if pivot["months"]:
        pivot_months = pivot["months"]
        # Header row
        ws3.cell(row=1, column=1, value="Categoria").fill = header_fill
        ws3.cell(row=1, column=1).font = header_font
        for ci, m in enumerate(pivot_months, 2):
            cell = ws3.cell(row=1, column=ci, value=m)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        total_h = ws3.cell(row=1, column=len(pivot_months) + 2, value="TOTAL")
        total_h.fill = header_fill
        total_h.font = header_font

        for ri, cat in enumerate(pivot["categories"], 2):
            ws3.cell(row=ri, column=1, value=cat["name"]).font = Font(bold=True)
            for ci, m in enumerate(pivot_months, 2):
                val = cat["values"].get(m, 0.0)
                cell = ws3.cell(row=ri, column=ci, value=val)
                cell.number_format = '#,##0.00'
            total_cat_cell = ws3.cell(row=ri, column=len(pivot_months) + 2, value=cat["total"])
            total_cat_cell.number_format = '#,##0.00'
            total_cat_cell.font = Font(bold=True)

        # Month totals row
        totals_row = len(pivot["categories"]) + 2
        ws3.cell(row=totals_row, column=1, value="TOTAL").font = Font(bold=True, color="FFFFFF")
        ws3.cell(row=totals_row, column=1).fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        for ci, m in enumerate(pivot_months, 2):
            cell = ws3.cell(row=totals_row, column=ci, value=pivot["month_totals"].get(m, 0.0))
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)
        grand_cell = ws3.cell(row=totals_row, column=len(pivot_months) + 2, value=pivot["grand_total"])
        grand_cell.number_format = '#,##0.00'
        grand_cell.font = Font(bold=True, color="FFFFFF")
        grand_cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

    return wb


@router.get("/excel")
def export_excel(
    months: int = 12,
    person: Optional[str] = None,
):
    df = get_transactions_df(months_back=months, person=person)
    wb = _build_workbook(df, months)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "dash_financeira_export.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
